import io
import os
from typing import Optional
import structlog

logger = structlog.get_logger()

async def extract_text_from_pdf(content: bytes) -> str:
    """Извлекает текст из PDF-файла."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return '\n'.join(text_parts)
    except Exception as exc:
        logger.warning('document_parser.pdf_error', error=str(exc))
        return ''

async def extract_text_from_docx(content: bytes) -> str:
    """Извлекает текст из DOCX."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        return '\n'.join([p.text for p in doc.paragraphs])
    except Exception as exc:
        logger.warning('document_parser.docx_error', error=str(exc))
        return ''

async def extract_text_from_xlsx(content: bytes) -> str:
    """Извлекает текст из XLSX."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = ' | '.join(str(cell) if cell is not None else '' for cell in row)
                if row_text.strip():
                    parts.append(row_text)
        return '\n'.join(parts)
    except Exception as exc:
        logger.warning('document_parser.xlsx_error', error=str(exc))
        return ''

async def extract_text(
    filename: str,
    content: bytes,
    mime_type: Optional[str] = None
) -> str:
    """Определяет тип файла и извлекает текст."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.pdf' or (mime_type and 'pdf' in mime_type):
        return await extract_text_from_pdf(content)
    elif ext == '.docx' or (mime_type and 'word' in mime_type):
        return await extract_text_from_docx(content)
    elif ext == '.xlsx' or (mime_type and 'spreadsheet' in mime_type):
        return await extract_text_from_xlsx(content)
    elif ext in ('.txt', '.csv'):
        return content.decode('utf-8', errors='ignore')
    else:
        logger.warning('document_parser.unsupported', filename=filename)
        return ''
